"""Pós-processa o master regenerado — ciclo de 12/10/2026.

1. "Samu Michel e Julia" -> vírgula faltando na origem (CPP1.2).
2. Sicredi: TUDO concluído (instrução do Caio, 12/10/2026).
3. Normalização de caixa dos status (app.py compara string exata).
4. Correção de prazo corrompido na origem.
5. Alocacao_Gantt derivado das Atividades (gantt.py, validado 125/125).
6. Carga_Esperada do PRIME refeita da aba do cronograma novo.
7. Abas novas: Eventos (Comunicação) e Monitorias (CPP 1.2).
"""
import os

# --- caminhos (sobrescreva por variavel de ambiente, se precisar) ---
#   CLEAR_CRONOGRAMAS  pasta com os Cronograma_*.xlsx  (default: ./cronogramas)
#   CLEAR_MASTER_ANTIGO  master vivo baixado do GitHub (default: ./CLEAR_Master_ANTIGO.xlsx)
#   CLEAR_OUT          master gerado  (default: /mnt/user-data/outputs/CLEAR_Master_2026.xlsx)
_AQUI = os.path.dirname(os.path.abspath(__file__))
Z = os.environ.get("CLEAR_CRONOGRAMAS", os.path.join(_AQUI, "cronogramas")).rstrip("/") + "/"
OLD = os.environ.get("CLEAR_MASTER_ANTIGO", os.path.join(_AQUI, "CLEAR_Master_ANTIGO.xlsx"))
OUT = os.environ.get("CLEAR_OUT", "/mnt/user-data/outputs/CLEAR_Master_2026.xlsx")

import datetime as dt
import openpyxl
import pandas as pd
import gantt
import parser as P   # parser.py deste repositorio


CANON = ["Bia B", "Bia S", "Caio", "Carol", "Cecilia", "Fabrícia", "Fred", "Hisrael",
         "Julia", "Junior", "Lorena", "Luan", "Luigi", "Lycia", "Michel", "Samu"]
NORM_PESSOA = {"Luiggi": "Luigi", "Zinho": "Michel", "Pleno": "Fred",
               "Senior": "Hisrael", "Sênior": "Hisrael", "Samuely": "Samu",
               "BiaS": "Bia S", "BiaB": "Bia B", "Júlia": "Julia", "Frederik": "Fred"}
STATUS_MAP = {"atrasado": "Atrasado", "pronto": "Concluído", "concluído": "Concluído",
              "concluido": "Concluído", "em andamento": "Em Andamento",
              "não iniciado": "Não Iniciado", "nao iniciado": "Não Iniciado",
              "entregável": "Entregável", "reunião": "Reunião",
              "realizado parcialmente": "Em Andamento"}

# prazos digitados errado na origem: (projeto, atividade) -> data correta
CORRECOES_PRAZO = {
    ("PRiME III", "Translate workbook to French"): dt.datetime(2027, 1, 22),  # vinha 207-01-22
}


def dmy(v):
    """Data que pode vir como datetime ou como texto dd/mm/aaaa."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
    d, _ = P.parse_prazo(v)
    return d


# ============================================================
# 1-4) correções nas Atividades
# ============================================================
wb = openpyxl.load_workbook(OUT)
ws = wb["Atividades"]
hdr = [c.value for c in ws[1]]
col = {n: hdr.index(n) + 1 for n in hdr if n}

n_comp = n_pz = 0
for r in range(2, ws.max_row + 1):
    raw = ws.cell(r, col["responsaveis_raw"]).value
    if isinstance(raw, str) and "Samu Michel" in raw:
        atual = [x.strip() for x in str(ws.cell(r, col["responsaveis"]).value or "").split(";") if x.strip()]
        novo = ["Samu", "Michel"] + [p for p in atual if p not in ("Samu", "Michel")]
        ws.cell(r, col["responsaveis"]).value = "; ".join(novo)
        ws.cell(r, col["n_responsaveis"]).value = len(novo)
        n_comp += 1
    chave = (ws.cell(r, col["projeto"]).value, str(ws.cell(r, col["atividade"]).value or "").strip())
    if chave in CORRECOES_PRAZO:
        ws.cell(r, col["prazo"]).value = CORRECOES_PRAZO[chave]
        n_pz += 1

n_sic = n_st = 0
for sh in ("Atividades", "Resp_Atividade"):
    w = wb[sh]
    h = [c.value for c in w[1]]
    ip, ist = h.index("projeto") + 1, h.index("status") + 1
    for r in range(2, w.max_row + 1):
        c = w.cell(r, ist)
        if w.cell(r, ip).value == "Sicredi":
            if c.value != "Concluído":
                c.value = "Concluído"
                n_sic += 1
            continue
        if isinstance(c.value, str):
            k = c.value.strip().lower()
            if k in STATUS_MAP and c.value != STATUS_MAP[k]:
                c.value = STATUS_MAP[k]
                n_st += 1
wb.save(OUT)
print(f"'Samu Michel': {n_comp} | prazos corrigidos: {n_pz} | Sicredi: {n_sic} | caixa status: {n_st}")

# ============================================================
# 5) Alocacao_Gantt
# ============================================================
ativ = pd.read_excel(OUT, sheet_name="Atividades")
g_old = pd.read_excel(OLD, sheet_name="Alocacao_Gantt")
g_new = gantt.derivar(ativ, g_old, CANON)

# ============================================================
# 6) Carga_Esperada (só o PRIME)
# ============================================================
ce = pd.read_excel(OLD, sheet_name="Carga_Esperada")
wsc = openpyxl.load_workbook(Z + "Cronograma_PRiME.xlsx", data_only=True)["Carga Esperada"]
rows = []
for r in range(5, 30):
    pes = wsc.cell(r, 1).value
    if not isinstance(pes, str):
        continue
    pes = NORM_PESSOA.get(pes.strip(), pes.strip())
    if pes not in CANON:
        continue
    for mi in range(2, 14):
        v = wsc.cell(r, mi).value
        if isinstance(v, str) and v.strip().upper() in {"B", "M", "A"}:
            rows.append(dict(pessoa=pes, projeto="PRIME", mes=mi - 1, nivel=v.strip().upper()))
ce_new = pd.concat([ce[ce["projeto"] != "PRIME"], pd.DataFrame(rows)], ignore_index=True)

# ============================================================
# 7a) Eventos — aba Eventos do Cronograma_Comunicação
# ============================================================
we = openpyxl.load_workbook(Z + "Cronograma_Comunicação_2026.xlsx", data_only=True)["Eventos"]
COLS_EV = ["categoria", "evento", "tipo", "projeto", "data", "data_fim", "formato",
           "local", "publico_alvo", "responsaveis", "status", "pecas_divulgacao", "obs"]
evs = []
for i, row in enumerate(we.iter_rows(values_only=True)):
    if i <= 3:                                     # titulo, instrucao, cabecalho, legenda
        continue
    ev = row[1]
    if not (isinstance(ev, str) and ev.strip()):
        continue
    if "EXEMPLO" in ev.upper():                    # linha de exemplo do template
        continue
    resp_list, _ = P.split_responsaveis(row[9])
    evs.append(dict(
        categoria=(str(row[0]).strip() if row[0] else None),
        evento=ev.strip(), tipo=(str(row[2]).strip() if row[2] else None),
        projeto=(str(row[3]).strip() if row[3] else None),
        data=dmy(row[4]), data_fim=dmy(row[5]),
        formato=(str(row[6]).strip() if row[6] else None),
        local=(str(row[7]).strip() if row[7] else None),
        publico_alvo=(str(row[8]).strip() if row[8] else None),
        responsaveis="; ".join(resp_list),
        status=(str(row[10]).strip() if row[10] else "A confirmar"),
        pecas_divulgacao=(str(row[11]).strip() if row[11] else None),
        obs=(str(row[12]).strip() if row[12] else None)))
df_ev = pd.DataFrame(evs, columns=COLS_EV)
df_ev.insert(0, "id_evento", range(1, len(df_ev) + 1))

# ============================================================
# 7b) Monitorias — aba Monitorias do CPP1.2 (sem a coluna de e-mail:
#     o repositorio e publico, e o painel nao precisa do contato)
# ============================================================
wm = openpyxl.load_workbook(Z + "Cronograma_IU_ CPP1.2.xlsx", data_only=True)["Monitorias"]
mons = []
for i, row in enumerate(wm.iter_rows(values_only=True)):
    if i <= 3:
        continue
    if not (isinstance(row[0], str) and row[0].strip()):
        continue
    mentor = str(row[3]).strip() if row[3] else None
    if mentor:
        mentor = NORM_PESSOA.get(mentor, mentor)
    mons.append(dict(
        monitoria=row[0].strip(), data=dmy(row[1]),
        horario=(str(row[2]).strip() if row[2] else None), mentor=mentor,
        uf_grupo=(str(row[5]).strip() if row[5] else None),
        turma=(str(row[6]).strip() if row[6] else None),
        status=(str(row[7]).strip() if row[7] else "Não Iniciado"),
        obs=(str(row[8]).strip() if row[8] else None)))
df_mon = pd.DataFrame(mons, columns=["monitoria", "data", "horario", "mentor",
                                     "uf_grupo", "turma", "status", "obs"])
df_mon.insert(0, "id_monitoria", range(1, len(df_mon) + 1))

# ============================================================
# grava
# ============================================================
wb = openpyxl.load_workbook(OUT)


def escreve(name, df, pos=None):
    if name in wb.sheetnames:
        pos = wb.sheetnames.index(name) if pos is None else pos
        del wb[name]
    w = wb.create_sheet(name, pos)
    w.append(list(df.columns))
    for rec in df.itertuples(index=False):
        w.append(["" if pd.isna(v) else v for v in rec])


escreve("Alocacao_Gantt", g_new)
escreve("Carga_Esperada", ce_new)
escreve("Eventos", df_ev, pos=3)        # logo depois de Resp_Atividade
escreve("Monitorias", df_mon, pos=4)

# README: registra as abas novas
ws = wb["README"]
linhas = {str(c.value).strip() for row in ws.iter_rows() for c in row if c.value}
if "Eventos" not in linhas:
    ws.append(["Eventos", "Eventos do centro (aba Eventos do cronograma de Comunicação)"])
    ws.append(["Monitorias", "Sessões de monitoria da Formação CPP 1.2, por mentor e grupo"])
wb.save(OUT)

print(f"Alocacao_Gantt: {len(g_old)} -> {len(g_new)} | {g_new['origem'].value_counts().to_dict()}")
print(f"Carga_Esperada: {len(ce)} -> {len(ce_new)}")
print(f"Eventos: {len(df_ev)} linhas | status {df_ev['status'].value_counts().to_dict()}")
print(f"Monitorias: {len(df_mon)} linhas | mentores {sorted(df_mon['mentor'].dropna().unique())}")
